import io
from pathlib import Path
from typing import BinaryIO, Dict, Iterable, List, Tuple

import PyPDF2
from pptx import Presentation


class NotesLoader:
    """
    Loads study notes from a folder or from uploaded files.
    Supports:
      - .txt, .md
      - .pdf (text extraction)
      - .pptx (slide text, tables, and speaker notes)
    """

    TEXT_EXTENSIONS = {".txt", ".md"}
    SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".pptx"}

    def load_notes_from_folder(self, folder: str) -> List[Dict[str, str]]:
        """
        Returns a list of dicts with:
          - 'name': file name
          - 'path': full path
          - 'content': extracted text
        """
        root = Path(folder).expanduser()

        if not root.exists():
            raise FileNotFoundError(f"Notes folder not found: {folder}")

        notes = []
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            content = self.extract_text(path.name, path.read_bytes())
            if content.strip():
                notes.append({"name": path.name, "path": str(path), "content": content})
        return notes

    def load_notes_from_uploads(
        self, files: Iterable[Tuple[str, bytes]]
    ) -> Tuple[List[Dict[str, str]], List[str]]:
        """
        files: iterable of (filename, raw bytes), e.g. from st.file_uploader.
        Returns (notes, unreadable_filenames).
        """
        notes: List[Dict[str, str]] = []
        unreadable: List[str] = []
        for name, data in files:
            content = self.extract_text(name, data)
            if content.strip():
                notes.append({"name": name, "path": name, "content": content})
            else:
                unreadable.append(name)
        return notes, unreadable

    def extract_text(self, filename: str, data: bytes) -> str:
        ext = Path(filename).suffix.lower()
        if ext in self.TEXT_EXTENSIONS:
            return data.decode("utf-8", errors="ignore")
        if ext == ".pdf":
            return self._read_pdf(io.BytesIO(data))
        if ext == ".pptx":
            return self._read_pptx(io.BytesIO(data))
        return ""

    def _read_pdf(self, stream: BinaryIO) -> str:
        text = ""
        try:
            reader = PyPDF2.PdfReader(stream)
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
        except Exception:
            return text
        return text

    def _read_pptx(self, stream: BinaryIO) -> str:
        parts: List[str] = []
        try:
            prs = Presentation(stream)
            for idx, slide in enumerate(prs.slides, start=1):
                slide_lines: List[str] = []
                for shape in slide.shapes:
                    slide_lines.extend(self._shape_text(shape))
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_lines.append(f"[Speaker notes] {notes}")
                if slide_lines:
                    parts.append(f"Slide {idx}:\n" + "\n".join(slide_lines))
        except Exception:
            return "\n\n".join(parts)
        return "\n\n".join(parts)

    def _shape_text(self, shape) -> List[str]:
        lines: List[str] = []
        if shape.shape_type == 6:  # group shape
            for child in shape.shapes:
                lines.extend(self._shape_text(child))
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                lines.append(text)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return lines
